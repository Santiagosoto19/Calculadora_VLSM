"""Exportador Excel al formato de clase VLSM (hoja tipo Tendencias).

Cada bloque ocupa las mismas filas a izquierda y derecha:
- A–H: título del grupo, encabezado y una fila por subred.
- I–AN: banner CALCULO, octetos y los 32 bits.
- AO–AQ: prefijo del contenedor, IP, prefijo hijo y nombre.
- AR–AU: tabla 2^n con fórmulas, una sola vez en la fila 2.

Colores de bits: naranja #FF9900 prefijo padre, cian #00FFFF bits de
subred, amarillo #FFFF00 bits de host.
"""

from __future__ import annotations

import ipaddress
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from core.models import ExportResult, IPv4Subnet, IPv6Result, MultiVLSMResult, VLSMResult, utc_now_iso

HEADER_BLUE = "2F5496"
BIT_PARENT = "FF9900"
BIT_SUBNET = "00FFFF"
BIT_HOST = "FFFF00"
WHITE = "FFFFFF"
INK = "000000"
THIN = Border(
    left=Side(style="thin", color="7F7F7F"),
    right=Side(style="thin", color="7F7F7F"),
    top=Side(style="thin", color="7F7F7F"),
    bottom=Side(style="thin", color="7F7F7F"),
)
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
LEFT = Alignment(horizontal="left", vertical="center")
CALIBRI = "Calibri"
BITS_COL = 9
RESULT_COL = 41
PREFIX_COL = 42
NAME_COL = 43
POW_COL = 44
TABLE_HEADERS = (
    "Subred",
    "#Hosts disponibles",
    "#Hosts usados",
    "IP",
    "Mascara de subred",
    "Primera direccion utilizable",
    "Ultima direccion utilizable",
    "Direccion Broadcast",
)


def build_workbook(result: ExportResult) -> Workbook:
    workbook = Workbook()
    sheet = workbook.active
    if isinstance(result, MultiVLSMResult):
        sheet.title = "VLSM Por EDIFICIO"
        _paint_ipv4_multi(sheet, result)
    elif isinstance(result, VLSMResult):
        sheet.title = "VLSM Por SUBRED"
        _paint_ipv4_classroom(sheet, result)
    else:
        sheet.title = "Subredes IPv6"
        _paint_ipv6_table(sheet, result)
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.paperSize = sheet.PAPERSIZE_A4
    sheet.page_setup.fitToPage = True
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    sheet.oddHeader.left.text = f"VLSM {result.base_cidr}"
    sheet.oddFooter.right.text = "Página &P de &N"
    return workbook


def workbook_to_bytes(workbook: Workbook) -> bytes:
    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def export_xlsx(result: ExportResult, path: str) -> str:
    build_workbook(result).save(path)
    return path


def _fill(color: str) -> PatternFill:
    return PatternFill("solid", fgColor=color)


def _font(*, bold: bool = False, color: str = INK, size: float = 11) -> Font:
    return Font(name=CALIBRI, bold=bold, color=color, size=size)


def _paint_ipv4_classroom(sheet: Worksheet, result: VLSMResult) -> None:
    parent = ipaddress.IPv4Network(result.base_cidr, strict=True)
    _write_powers_table(sheet)
    _paint_aligned_block(
        sheet,
        1,
        banner=f"CALCULO {result.base_cidr}",
        parent=parent,
        subnets=list(result.subnets),
        left_title=result.base_cidr,
    )
    _set_classroom_widths(sheet)
    sheet.sheet_view.showGridLines = True


def _paint_ipv4_multi(sheet: Worksheet, result: MultiVLSMResult) -> None:
    parent = ipaddress.IPv4Network(result.base_cidr, strict=True)
    _write_powers_table(sheet)
    row = 1
    if result.containers:
        row = _paint_aligned_block(
            sheet,
            row,
            banner=f"CALCULO {result.base_cidr}",
            parent=parent,
            subnets=list(result.containers),
            left_title=None,
        )
    for group in result.groups:
        inner = ipaddress.IPv4Network(group.container.cidr, strict=True)
        row = _paint_aligned_block(
            sheet,
            row,
            banner=f"CALCULO {group.name}",
            parent=inner,
            subnets=list(group.result.subnets),
            left_title=group.name,
        )
    _set_classroom_widths(sheet)
    sheet.sheet_view.showGridLines = True


def _paint_aligned_block(
    sheet: Worksheet,
    start_row: int,
    *,
    banner: str,
    parent: ipaddress.IPv4Network,
    subnets: list[IPv4Subnet],
    left_title: str | None,
) -> int:
    """Tabla A–H y rejilla I–AQ del mismo bloque, empezando en start_row."""
    _write_banner(sheet, start_row, banner, BITS_COL, BITS_COL + 31)
    _write_binary_header(sheet, parent, row=start_row + 1)
    _write_bit_row(
        sheet,
        start_row + 2,
        int(parent.network_address),
        parent.prefixlen,
        parent.prefixlen,
    )

    row = start_row + 3
    parent_end = int(parent.broadcast_address)
    for index, subnet in enumerate(subnets, start=1):
        row = _append_calculo(
            sheet,
            row,
            index,
            subnet,
            parent_prefix=parent.prefixlen,
            parent_end=parent_end,
        )

    table_end = start_row
    if left_title is not None:
        table_end = _write_group_table(sheet, start_row, left_title, subnets)
    last = max(row - 1, table_end)
    return last + 2


def _write_group_table(sheet: Worksheet, start_row: int, title: str, subnets: list[IPv4Subnet]) -> int:
    sheet.merge_cells(start_row=start_row, start_column=1, end_row=start_row, end_column=8)
    cell = sheet.cell(start_row, 1, title)
    cell.font = _font()
    cell.alignment = CENTER

    header_row = start_row + 1
    for col, header in enumerate(TABLE_HEADERS, start=1):
        head = sheet.cell(header_row, col, header)
        head.fill = _fill(HEADER_BLUE)
        head.font = _font(color=WHITE)
        head.alignment = CENTER
        head.border = THIN
    sheet.row_dimensions[header_row].height = 32

    row = header_row
    for subnet in subnets:
        row += 1
        values: list[object] = [
            subnet.name,
            subnet.assigned_usable,
            subnet.requested_hosts,
            subnet.network,
            str(subnet.netmask),
            subnet.first_usable,
            subnet.last_usable,
            subnet.broadcast,
        ]
        for col, value in enumerate(values, start=1):
            item = sheet.cell(row, col, value)
            item.font = _font()
            item.border = THIN
            item.alignment = CENTER if col > 1 else LEFT
            if col in {2, 3}:
                item.number_format = "0"
            elif col == 5:
                item.number_format = "@"
    return row


def _append_calculo(
    sheet: Worksheet,
    calc_row: int,
    index: int,
    subnet: IPv4Subnet,
    *,
    parent_prefix: int,
    parent_end: int,
) -> int:
    child = ipaddress.IPv4Network(subnet.cidr, strict=True)
    _write_banner(sheet, calc_row, f"CALCULO {index}", BITS_COL, BITS_COL + 31)

    assigned_row = calc_row + 1
    _write_bit_row(sheet, assigned_row, int(child.network_address), parent_prefix, child.prefixlen)
    _write_note(sheet, RESULT_COL, assigned_row, subnet.network)
    _write_note(sheet, PREFIX_COL, assigned_row, f"/{subnet.prefixlen}")
    _write_note(sheet, NAME_COL, assigned_row, subnet.name)

    next_addr = int(child.network_address) + subnet.block_size
    if next_addr <= parent_end:
        next_row = assigned_row + 1
        _write_bit_row(sheet, next_row, next_addr, parent_prefix, child.prefixlen)
        _write_note(sheet, RESULT_COL, next_row, str(ipaddress.IPv4Address(next_addr)))
        _write_note(sheet, PREFIX_COL, next_row, f"/{child.prefixlen}")
        _write_note(sheet, NAME_COL, next_row, "Siguiente calculo")
        return next_row + 1
    return assigned_row + 1


def _write_binary_header(sheet: Worksheet, parent: ipaddress.IPv4Network, *, row: int) -> None:
    octets = str(parent.network_address).split(".")
    for group, octet in enumerate(octets):
        start = BITS_COL + group * 8
        end = start + 7
        sheet.merge_cells(start_row=row, start_column=start, end_row=row, end_column=end)
        cell = sheet.cell(row, start, f"{octet}.")
        cell.font = _font()
        cell.alignment = CENTER
        cell.number_format = "@"
    prefix = sheet.cell(row, RESULT_COL, f"/{parent.prefixlen}")
    prefix.font = _font()
    prefix.alignment = CENTER


def _write_bit_row(
    sheet: Worksheet,
    row: int,
    address: int,
    parent_prefix: int,
    child_prefix: int,
) -> None:
    for index in range(32):
        col = BITS_COL + index
        bit = (address >> (31 - index)) & 1
        cell = sheet.cell(row, col, bit)
        if index < parent_prefix:
            color = BIT_PARENT
        elif index < child_prefix:
            color = BIT_SUBNET
        else:
            color = BIT_HOST
        cell.fill = _fill(color)
        cell.font = _font()
        cell.alignment = CENTER
        cell.number_format = "0"


def _write_banner(sheet: Worksheet, row: int, text: str, start_col: int, end_col: int) -> None:
    sheet.merge_cells(start_row=row, start_column=start_col, end_row=row, end_column=end_col)
    cell = sheet.cell(row, start_col, text)
    cell.font = _font()
    cell.alignment = CENTER


def _write_note(sheet: Worksheet, col: int, row: int, value: object) -> None:
    cell = sheet.cell(row, col, value)
    cell.font = _font()
    cell.alignment = CENTER


def _write_powers_table(sheet: Worksheet) -> None:
    headers = ("Base", "Bit prestado", "#Subredes", "#Hosts")
    for offset, header in enumerate(headers):
        cell = sheet.cell(2, POW_COL + offset, header)
        cell.font = _font()
        cell.alignment = CENTER

    base_letter = get_column_letter(POW_COL)
    bit_letter = get_column_letter(POW_COL + 1)
    power_letter = get_column_letter(POW_COL + 2)
    for n in range(16):
        row = 3 + n
        base = sheet.cell(row, POW_COL, 2)
        borrowed = sheet.cell(row, POW_COL + 1, n)
        power = sheet.cell(row, POW_COL + 2, f"={base_letter}{row}^{bit_letter}{row}")
        if n < 2:
            hosts = sheet.cell(row, POW_COL + 3, 0)
            hosts.number_format = "0"
        else:
            hosts = sheet.cell(row, POW_COL + 3, f"={power_letter}{row}-{base_letter}{row}")
        for cell in (base, borrowed, power, hosts):
            cell.font = _font()
            cell.alignment = CENTER
        base.number_format = "0"
        borrowed.number_format = "0"


def _set_classroom_widths(sheet: Worksheet) -> None:
    widths = {
        1: 20,
        2: 9.38,
        3: 12.13,
        4: 16.13,
        5: 17.38,
        6: 21.25,
        7: 20.25,
        8: 20,
        RESULT_COL: 13.5,
        PREFIX_COL: 9.38,
        NAME_COL: 17.13,
        POW_COL: 4.5,
        POW_COL + 1: 10.25,
        POW_COL + 2: 9,
        POW_COL + 3: 7.25,
    }
    for col, width in widths.items():
        sheet.column_dimensions[get_column_letter(col)].width = width
    for col in range(BITS_COL, BITS_COL + 32):
        sheet.column_dimensions[get_column_letter(col)].width = 2


def _paint_ipv6_table(sheet: Worksheet, result: IPv6Result) -> None:
    sheet.merge_cells("A1:H1")
    title = sheet.cell(1, 1, f"IPv6  {result.base_cidr}")
    title.fill = _fill(HEADER_BLUE)
    title.font = _font(bold=True, color=WHITE, size=14)
    title.alignment = CENTER
    for col in range(1, 9):
        sheet.cell(1, col).fill = _fill(HEADER_BLUE)
        sheet.cell(1, col).border = THIN

    headers = [
        "Subred",
        "Dispositivos",
        "Prefijo",
        "CIDR comprimido",
        "Expandido",
        "Primera asignable",
        "Ultima asignable",
        "Anycast router/subred",
    ]
    for col, header in enumerate(headers, start=1):
        cell = sheet.cell(3, col, header)
        cell.fill = _fill(HEADER_BLUE)
        cell.font = _font(color=WHITE)
        cell.alignment = CENTER
        cell.border = THIN
    sheet.row_dimensions[3].height = 28

    for offset, subnet in enumerate(result.subnets):
        row = 4 + offset
        values = [
            subnet.name,
            subnet.requested_devices,
            f"/{subnet.prefixlen}",
            subnet.compressed,
            subnet.expanded,
            subnet.first_assignable,
            subnet.last_assignable,
            subnet.router_anycast,
        ]
        for col, value in enumerate(values, start=1):
            cell = sheet.cell(row, col, value)
            cell.font = _font()
            cell.border = THIN
            cell.alignment = LEFT if col in {1, 4, 5} else CENTER

    leftover_row = 5 + len(result.subnets)
    sheet.cell(leftover_row, 1, "Sobrante").font = _font(bold=True)
    sheet.cell(leftover_row, 2, ", ".join(result.leftover_cidrs[:12]) or "Ninguno")
    for col in range(1, 9):
        sheet.column_dimensions[get_column_letter(col)].width = 28
    sheet.column_dimensions["A"].width = 22
    sheet.freeze_panes = "A4"


def stamp_generated(result: ExportResult) -> str:
    return result.generated_at or utc_now_iso()
